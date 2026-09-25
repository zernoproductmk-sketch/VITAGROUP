from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook
from sqlalchemy import text

from .config import settings
from .database import engine
from .integrations.yandex_disk import (
    YandexDiskClient,
    configured_plan_source,
    public_key_hash,
)
from .printed_1c_plan import (
    detect_printed_1c_form,
    inspect_printed_1c_form,
    parse_printed_1c_form,
)


FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "erp_guid": (
        "erp_guid", "guid", "идентификатор", "идентификатор erp",
    ),
    "task_id": (
        "id задания", "№ задания", "номер задания",
        "задание на производство №", "задание на производство",
    ),
    "business_date": (
        "дата", "дата выполнения", "плановая дата",
        "дата производства", "дата смены",
    ),
    "shift_code": ("смена",),
    "workshop": ("цех", "подразделение"),
    "equipment_code": (
        "станок", "линия", "оборудование", "рабочий центр",
    ),
    "order_no": (
        "№ заказа", "номер заказа", "заказ",
        "заказ на производство", "заказ производства",
    ),
    "article": (
        "артикул", "код артикула", "номенклатура",
        "код номенклатуры",
    ),
    "customer": ("клиент", "контрагент", "заказчик"),
    "tech_card": (
        "техкарта", "технологическая карта",
        "спецификация", "ресурсная спецификация",
    ),
    "plan_qty_pcs": (
        "план, шт", "план шт", "количество, шт",
        "количество шт", "плановое количество, шт",
    ),
    "pcs_per_box": (
        "шт/короб", "шт в коробе", "штук в коробе",
        "количество в коробе",
    ),
    "boxes_per_pallet": (
        "коробов/паллет", "коробов на паллете",
        "коробов в паллете",
    ),
    "plan_kg": (
        "план, кг", "план кг", "количество, кг",
        "количество кг", "плановое количество, кг",
    ),
    "source_status": ("статус",),
}

CORE_FIELDS = {"business_date", "order_no", "article"}


def _clean_header(value: Any) -> str:
    text_value = str(value or "").strip().lower()
    text_value = text_value.replace("ё", "е")
    text_value = re.sub(r"\s+", " ", text_value)
    return text_value


def _build_mapping(headers: list[Any]) -> dict[str, int]:
    normalized = [_clean_header(value) for value in headers]
    mapping: dict[str, int] = {}

    for field_name, aliases in FIELD_ALIASES.items():
        alias_set = {_clean_header(alias) for alias in aliases}
        for index, header in enumerate(normalized):
            if header in alias_set:
                mapping[field_name] = index
                break

    return mapping


def _to_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text_value = str(value).strip().replace(" ", "").replace(",", ".")
    try:
        return Decimal(text_value)
    except InvalidOperation:
        return None


def _to_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            return date(1899, 12, 30) + __import__("datetime").timedelta(days=int(value))
        except Exception:
            return None

    text_value = str(value).strip()
    for fmt in (
        "%d.%m.%Y",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d.%m.%y",
    ):
        try:
            return datetime.strptime(text_value, fmt).date()
        except ValueError:
            pass
    return None


def _normalize_shift(value: Any) -> str | None:
    text_value = str(value or "").strip().upper()
    if text_value in {"DAY", "ДЕНЬ", "ДНЕВНАЯ", "1"}:
        return "DAY"
    if text_value in {"NIGHT", "НОЧЬ", "НОЧНАЯ", "2"}:
        return "NIGHT"
    return None


def _read_xlsx(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(
        io.BytesIO(content),
        read_only=True,
        data_only=True,
    )
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        sheets.append({"name": worksheet.title, "rows": rows})
    return sheets


def _read_xls(content: bytes) -> list[dict[str, Any]]:
    workbook = xlrd.open_workbook(file_contents=content)
    sheets: list[dict[str, Any]] = []
    for sheet in workbook.sheets():
        rows = [sheet.row_values(index) for index in range(sheet.nrows)]
        sheets.append({"name": sheet.name, "rows": rows})
    return sheets


def _read_csv(content: bytes) -> list[dict[str, Any]]:
    decoded = None
    for encoding in ("utf-8-sig", "cp1251", "utf-8"):
        try:
            decoded = content.decode(encoding)
            break
        except UnicodeDecodeError:
            pass
    if decoded is None:
        raise ValueError("CSV encoding is not supported")

    sample = decoded[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t,")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"

    rows = list(csv.reader(io.StringIO(decoded), dialect))
    return [{"name": "CSV", "rows": rows}]


def parse_workbook(content: bytes, file_name: str) -> list[dict[str, Any]]:
    suffix = Path(file_name).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(content)
    if suffix == ".xls":
        return _read_xls(content)
    if suffix in {".csv", ".txt"}:
        return _read_csv(content)
    raise ValueError(
        f"Unsupported plan file format: {suffix or 'without extension'}"
    )


def _find_header_row(rows: list[list[Any]], scan_limit: int = 30) -> tuple[int, list[Any], dict[str, int]]:
    best: tuple[int, list[Any], dict[str, int]] | None = None
    best_score = -1

    for index, row in enumerate(rows[:scan_limit]):
        mapping = _build_mapping(row)
        score = len(mapping)
        if score > best_score:
            best = (index, row, mapping)
            best_score = score

    if best is None:
        return 0, [], {}
    return best


def inspect_workbook(content: bytes, file_name: str) -> dict[str, Any]:
    sheets = parse_workbook(content, file_name)
    result = []

    for sheet in sheets:
        if detect_printed_1c_form(sheet["rows"]):
            printed = inspect_printed_1c_form(sheet["rows"])
            result.append(
                {
                    "name": sheet["name"],
                    "schema_type": printed["schema_type"],
                    "header_row": None,
                    "headers": [],
                    "mapping": {},
                    "recognized_fields": [
                        "task_id",
                        "order_no",
                        "business_date",
                        "workshop",
                        "article",
                        "product_name",
                        "tech_card",
                        "plan_qty_pcs",
                        "route_operation",
                        "route_equipment_name",
                        "norm_hours",
                        "ideal_rate_per_hour",
                    ],
                    "core_fields_found": [
                        field
                        for field in ("business_date", "order_no", "article")
                        if (
                            field != "article"
                            or any(item.get("article") for item in printed.get("outputs", []))
                        )
                        and (
                            field != "business_date"
                            or printed.get("business_date")
                        )
                        and (
                            field != "order_no"
                            or printed.get("order_no")
                        )
                    ],
                    "printed_form": printed,
                    "sample_rows": [],
                }
            )
            continue

        header_row, headers, mapping = _find_header_row(sheet["rows"])
        sample_rows = []
        for row in sheet["rows"][header_row + 1: header_row + 6]:
            if any(value not in (None, "") for value in row):
                sample_rows.append(
                    [None if value is None else str(value) for value in row]
                )

        result.append(
            {
                "name": sheet["name"],
                "header_row": header_row + 1,
                "headers": [None if value is None else str(value) for value in headers],
                "mapping": mapping,
                "recognized_fields": sorted(mapping.keys()),
                "core_fields_found": sorted(CORE_FIELDS.intersection(mapping.keys())),
                "sample_rows": sample_rows,
            }
        )

    return {"sheets": result}


def _choose_sheet(inspected: dict[str, Any]) -> dict[str, Any] | None:
    sheets = inspected.get("sheets") or []
    if not sheets:
        return None
    return max(
        sheets,
        key=lambda item: (
            1 if item.get("schema_type") == "1C_PRINTED_PRODUCTION_TASK" else 0,
            len(item.get("core_fields_found") or []),
            len(item.get("recognized_fields") or []),
        ),
    )


def _cell(row: list[Any], mapping: dict[str, int], field_name: str):
    index = mapping.get(field_name)
    if index is None or index >= len(row):
        return None
    return row[index]


def _normalized_row(
    row: list[Any],
    mapping: dict[str, int],
) -> dict[str, Any]:
    return {
        "erp_guid": _cell(row, mapping, "erp_guid"),
        "task_id": _cell(row, mapping, "task_id"),
        "business_date": _to_date(_cell(row, mapping, "business_date")),
        "shift_code": _normalize_shift(_cell(row, mapping, "shift_code")),
        "workshop": _cell(row, mapping, "workshop"),
        "equipment_code": _cell(row, mapping, "equipment_code"),
        "order_no": _cell(row, mapping, "order_no"),
        "article": _cell(row, mapping, "article"),
        "customer": _cell(row, mapping, "customer"),
        "tech_card": _cell(row, mapping, "tech_card"),
        "plan_qty_pcs": _to_decimal(_cell(row, mapping, "plan_qty_pcs")),
        "pcs_per_box": _to_decimal(_cell(row, mapping, "pcs_per_box")),
        "boxes_per_pallet": _to_decimal(_cell(row, mapping, "boxes_per_pallet")),
        "plan_kg": _to_decimal(_cell(row, mapping, "plan_kg")),
        "source_status": _cell(row, mapping, "source_status"),
    }


def _apply_missing_shift_policy(item: dict[str, Any]) -> dict[str, Any]:
    if item.get("shift_code"):
        return item

    policy = settings.erp_plan_missing_shift_policy.strip().upper()
    if policy != "NEXT_DAY":
        return item

    business_date = item.get("business_date")
    if business_date:
        item["business_date"] = business_date + timedelta(days=1)

    shift_code = settings.erp_plan_default_shift_code.strip().upper() or "DAY"
    if shift_code not in {"DAY", "NIGHT"}:
        raise ValueError(
            "ERP_PLAN_DEFAULT_SHIFT_CODE must be DAY or NIGHT"
        )
    item["shift_code"] = shift_code
    return item


def _row_status(item: dict[str, Any]) -> tuple[str, str | None]:
    missing = []
    for field_name in ("business_date", "order_no", "article"):
        if item.get(field_name) in (None, ""):
            missing.append(field_name)

    if not missing:
        return "VALID", None

    if len(missing) == 3:
        return "IGNORED", "No core ERP plan fields"

    return "PARTIAL", f"Missing fields: {', '.join(missing)}"


async def yandex_plan_preview() -> dict[str, Any]:
    public_url, resource_path = configured_plan_source()
    client = YandexDiskClient()
    try:
        metadata = await client.metadata(public_url, resource_path)
        if metadata.get("type") != "file":
            return {
                "configured": True,
                "resource_type": metadata.get("type"),
                "name": metadata.get("name"),
                "message": "Configured Yandex resource is a folder; specify YANDEX_PLAN_RESOURCE_PATH",
                "items": [
                    {
                        "name": item.get("name"),
                        "type": item.get("type"),
                        "path": item.get("path"),
                        "size": item.get("size"),
                        "modified": item.get("modified"),
                    }
                    for item in ((metadata.get("_embedded") or {}).get("items") or [])[:100]
                ],
            }

        file_name = metadata.get("name") or "plan.xlsx"
        content = await client.download_bytes(public_url, resource_path)
        inspection = inspect_workbook(content, file_name)

        return {
            "configured": True,
            "resource_type": "file",
            "name": file_name,
            "size": metadata.get("size"),
            "modified": metadata.get("modified"),
            **inspection,
        }
    finally:
        await client.close()


def _persist_printed_1c_plan(
    *,
    parsed: dict[str, Any],
    file_name: str,
    sheet_name: str,
    metadata: dict[str, Any],
    source_hash: str,
) -> dict[str, Any]:
    mapping = {
        "schema_type": "1C_PRINTED_PRODUCTION_TASK",
        "source": "label_and_section_parser",
    }

    with engine.begin() as connection:
        batch_id = connection.execute(
            text(
                """
                INSERT INTO erp_plan_import_batches (
                    source_file_name,
                    source_file_modified,
                    source_file_size,
                    source_public_key_hash,
                    source_sheet,
                    mapping
                ) VALUES (
                    :file_name,
                    :modified,
                    :size,
                    :source_hash,
                    :sheet_name,
                    CAST(:mapping AS jsonb)
                )
                RETURNING id
                """
            ),
            {
                "file_name": file_name,
                "modified": metadata.get("modified"),
                "size": metadata.get("size"),
                "source_hash": source_hash,
                "sheet_name": sheet_name,
                "mapping": json.dumps(mapping, ensure_ascii=False),
            },
        ).scalar_one()

        applied = skipped = errors = 0

        for item in parsed["outputs"]:
            item = _apply_missing_shift_policy(dict(item))
            status, error_message = _row_status(item)
            if status == "IGNORED":
                skipped += 1
            else:
                applied += 1

            source_row = int(item.get("source_row") or 1)
            source_record_key = "|".join(
                [
                    "YANDEX_DISK",
                    source_hash,
                    file_name,
                    sheet_name,
                    str(source_row),
                    str(item.get("task_id") or item.get("order_no") or ""),
                    str(item.get("article") or ""),
                ]
            )

            raw_data = {
                "schema_type": parsed["schema_type"],
                "task_id": parsed.get("task_id"),
                "order_no": parsed.get("order_no"),
                "business_date": parsed.get("business_date"),
                "shift_code": parsed.get("shift_code"),
                "execution_at": parsed.get("execution_at"),
                "workshop": parsed.get("workshop"),
                "product_name": item.get("product_name"),
                "specification": parsed.get("specification"),
                "route_operations": item.get("route_operations") or [],
            }

            try:
                connection.execute(
                    text(
                        """
                        INSERT INTO erp_plan_staging (
                            import_batch_id,
                            source_file_name,
                            source_sheet,
                            source_row,
                            source_record_key,
                            erp_guid,
                            task_id,
                            business_date,
                            shift_code,
                            workshop,
                            equipment_code,
                            order_no,
                            article,
                            customer,
                            tech_card,
                            plan_qty_pcs,
                            pcs_per_box,
                            boxes_per_pallet,
                            plan_kg,
                            source_status,
                            product_name,
                            output_unit,
                            route_operation,
                            route_equipment_name,
                            route_equipment_hint,
                            norm_hours,
                            ideal_rate_per_hour,
                            route_operations,
                            raw_data,
                            row_status,
                            error_message
                        ) VALUES (
                            :import_batch_id,
                            :source_file_name,
                            :source_sheet,
                            :source_row,
                            :source_record_key,
                            :erp_guid,
                            :task_id,
                            :business_date,
                            :shift_code,
                            :workshop,
                            :equipment_code,
                            :order_no,
                            :article,
                            :customer,
                            :tech_card,
                            :plan_qty_pcs,
                            :pcs_per_box,
                            :boxes_per_pallet,
                            :plan_kg,
                            :source_status,
                            :product_name,
                            :output_unit,
                            :route_operation,
                            :route_equipment_name,
                            :route_equipment_hint,
                            :norm_hours,
                            :ideal_rate_per_hour,
                            CAST(:route_operations AS jsonb),
                            CAST(:raw_data AS jsonb),
                            :row_status,
                            :error_message
                        )
                        ON CONFLICT (source_record_key)
                        DO UPDATE SET
                            import_batch_id = EXCLUDED.import_batch_id,
                            business_date = EXCLUDED.business_date,
                            workshop = EXCLUDED.workshop,
                            order_no = EXCLUDED.order_no,
                            article = EXCLUDED.article,
                            tech_card = EXCLUDED.tech_card,
                            plan_qty_pcs = EXCLUDED.plan_qty_pcs,
                            plan_kg = EXCLUDED.plan_kg,
                            product_name = EXCLUDED.product_name,
                            output_unit = EXCLUDED.output_unit,
                            route_operation = EXCLUDED.route_operation,
                            route_equipment_name = EXCLUDED.route_equipment_name,
                            route_equipment_hint = EXCLUDED.route_equipment_hint,
                            norm_hours = EXCLUDED.norm_hours,
                            ideal_rate_per_hour = EXCLUDED.ideal_rate_per_hour,
                            route_operations = EXCLUDED.route_operations,
                            raw_data = EXCLUDED.raw_data,
                            row_status = EXCLUDED.row_status,
                            error_message = EXCLUDED.error_message,
                            updated_at = now()
                        """
                    ),
                    {
                        "import_batch_id": batch_id,
                        "source_file_name": file_name,
                        "source_sheet": sheet_name,
                        "source_row": source_row,
                        "source_record_key": source_record_key,
                        **{
                            key: item.get(key)
                            for key in (
                                "erp_guid",
                                "task_id",
                                "business_date",
                                "shift_code",
                                "workshop",
                                "equipment_code",
                                "order_no",
                                "article",
                                "customer",
                                "tech_card",
                                "plan_qty_pcs",
                                "pcs_per_box",
                                "boxes_per_pallet",
                                "plan_kg",
                                "source_status",
                                "product_name",
                                "output_unit",
                                "route_operation",
                                "route_equipment_name",
                                "route_equipment_hint",
                                "norm_hours",
                                "ideal_rate_per_hour",
                            )
                        },
                        "route_operations": json.dumps(
                            item.get("route_operations") or [],
                            ensure_ascii=False,
                            default=str,
                        ),
                        "raw_data": json.dumps(
                            raw_data,
                            ensure_ascii=False,
                            default=str,
                        ),
                        "row_status": status,
                        "error_message": error_message,
                    },
                )
            except Exception:
                errors += 1
                applied = max(0, applied - 1)

        batch_status = "COMPLETED_WITH_ERRORS" if errors else "COMPLETED"
        connection.execute(
            text(
                """
                UPDATE erp_plan_import_batches
                SET status = :status,
                    rows_read = :rows_read,
                    rows_applied = :rows_applied,
                    rows_skipped = :rows_skipped,
                    rows_error = :rows_error,
                    completed_at = now()
                WHERE id = :batch_id
                """
            ),
            {
                "status": batch_status,
                "rows_read": len(parsed["outputs"]),
                "rows_applied": applied,
                "rows_skipped": skipped,
                "rows_error": errors,
                "batch_id": batch_id,
            },
        )

    return {
        "status": batch_status,
        "batch_id": str(batch_id),
        "file_name": file_name,
        "sheet": sheet_name,
        "schema_type": parsed["schema_type"],
        "task_id": parsed.get("task_id"),
        "order_no": parsed.get("order_no"),
        "business_date": (
            parsed["business_date"].isoformat()
            if parsed.get("business_date")
            else None
        ),
        "outputs": len(parsed["outputs"]),
        "rows_applied": applied,
        "rows_skipped": skipped,
        "rows_error": errors,
    }


async def import_yandex_plan() -> dict[str, Any]:
    public_url, resource_path = configured_plan_source()
    client = YandexDiskClient()
    try:
        metadata = await client.metadata(public_url, resource_path)
        if metadata.get("type") != "file":
            raise ValueError(
                "YANDEX_PLAN_PUBLIC_URL points to a folder; configure YANDEX_PLAN_RESOURCE_PATH"
            )

        file_name = metadata.get("name") or "plan.xlsx"
        content = await client.download_bytes(public_url, resource_path)
    finally:
        await client.close()

    workbook = parse_workbook(content, file_name)
    inspection = inspect_workbook(content, file_name)
    selected = _choose_sheet(inspection)
    if not selected:
        raise ValueError("Workbook does not contain readable sheets")

    sheet_name = selected["name"]

    if selected.get("schema_type") == "1C_PRINTED_PRODUCTION_TASK":
        sheet = next(item for item in workbook if item["name"] == sheet_name)
        parsed = parse_printed_1c_form(sheet["rows"])
        return _persist_printed_1c_plan(
            parsed=parsed,
            file_name=file_name,
            sheet_name=sheet_name,
            metadata=metadata,
            source_hash=public_key_hash(public_url),
        )

    core_found = set(selected.get("core_fields_found") or [])
    if len(core_found) < 2:
        raise ValueError(
            "Plan file schema is not recognized: at least two core fields "
            "(date, order, article) are required"
        )

    header_row_index = int(selected["header_row"]) - 1
    mapping = selected["mapping"]
    sheet = next(item for item in workbook if item["name"] == sheet_name)

    modified = metadata.get("modified")
    source_hash = public_key_hash(public_url)

    with engine.begin() as connection:
        batch_id = connection.execute(
            text(
                """
                INSERT INTO erp_plan_import_batches (
                    source_file_name,
                    source_file_modified,
                    source_file_size,
                    source_public_key_hash,
                    source_sheet,
                    mapping
                ) VALUES (
                    :file_name,
                    :modified,
                    :size,
                    :source_hash,
                    :sheet_name,
                    CAST(:mapping AS jsonb)
                )
                RETURNING id
                """
            ),
            {
                "file_name": file_name,
                "modified": modified,
                "size": metadata.get("size"),
                "source_hash": source_hash,
                "sheet_name": sheet_name,
                "mapping": json.dumps(mapping, ensure_ascii=False),
            },
        ).scalar_one()

        read = applied = skipped = errors = 0

        for source_row, row in enumerate(
            sheet["rows"][header_row_index + 1:],
            start=header_row_index + 2,
        ):
            if not any(value not in (None, "") for value in row):
                continue

            read += 1
            raw_data = {
                str(index + 1): value
                for index, value in enumerate(row)
                if value not in (None, "")
            }

            try:
                item = _apply_missing_shift_policy(_normalized_row(row, mapping))
                status, error_message = _row_status(item)

                if status == "IGNORED":
                    skipped += 1
                else:
                    applied += 1

                source_record_key = "|".join(
                    [
                        "YANDEX_DISK",
                        source_hash,
                        file_name,
                        sheet_name,
                        str(source_row),
                        str(item.get("erp_guid") or item.get("task_id") or item.get("order_no") or ""),
                    ]
                )

                connection.execute(
                    text(
                        """
                        INSERT INTO erp_plan_staging (
                            import_batch_id,
                            source_file_name,
                            source_sheet,
                            source_row,
                            source_record_key,
                            erp_guid,
                            task_id,
                            business_date,
                            shift_code,
                            workshop,
                            equipment_code,
                            order_no,
                            article,
                            customer,
                            tech_card,
                            plan_qty_pcs,
                            pcs_per_box,
                            boxes_per_pallet,
                            plan_kg,
                            source_status,
                            raw_data,
                            row_status,
                            error_message
                        ) VALUES (
                            :import_batch_id,
                            :source_file_name,
                            :source_sheet,
                            :source_row,
                            :source_record_key,
                            :erp_guid,
                            :task_id,
                            :business_date,
                            :shift_code,
                            :workshop,
                            :equipment_code,
                            :order_no,
                            :article,
                            :customer,
                            :tech_card,
                            :plan_qty_pcs,
                            :pcs_per_box,
                            :boxes_per_pallet,
                            :plan_kg,
                            :source_status,
                            CAST(:raw_data AS jsonb),
                            :row_status,
                            :error_message
                        )
                        ON CONFLICT (source_record_key)
                        DO UPDATE SET
                            import_batch_id = EXCLUDED.import_batch_id,
                            business_date = EXCLUDED.business_date,
                            shift_code = EXCLUDED.shift_code,
                            workshop = EXCLUDED.workshop,
                            equipment_code = EXCLUDED.equipment_code,
                            order_no = EXCLUDED.order_no,
                            article = EXCLUDED.article,
                            customer = EXCLUDED.customer,
                            tech_card = EXCLUDED.tech_card,
                            plan_qty_pcs = EXCLUDED.plan_qty_pcs,
                            pcs_per_box = EXCLUDED.pcs_per_box,
                            boxes_per_pallet = EXCLUDED.boxes_per_pallet,
                            plan_kg = EXCLUDED.plan_kg,
                            source_status = EXCLUDED.source_status,
                            raw_data = EXCLUDED.raw_data,
                            row_status = EXCLUDED.row_status,
                            error_message = EXCLUDED.error_message,
                            updated_at = now()
                        """
                    ),
                    {
                        "import_batch_id": batch_id,
                        "source_file_name": file_name,
                        "source_sheet": sheet_name,
                        "source_row": source_row,
                        "source_record_key": source_record_key,
                        **item,
                        "raw_data": json.dumps(raw_data, ensure_ascii=False, default=str),
                        "row_status": status,
                        "error_message": error_message,
                    },
                )

            except Exception as exc:
                errors += 1

        batch_status = "COMPLETED_WITH_ERRORS" if errors else "COMPLETED"
        connection.execute(
            text(
                """
                UPDATE erp_plan_import_batches
                SET status = :status,
                    rows_read = :rows_read,
                    rows_applied = :rows_applied,
                    rows_skipped = :rows_skipped,
                    rows_error = :rows_error,
                    completed_at = now()
                WHERE id = :batch_id
                """
            ),
            {
                "status": batch_status,
                "rows_read": read,
                "rows_applied": applied,
                "rows_skipped": skipped,
                "rows_error": errors,
                "batch_id": batch_id,
            },
        )

    return {
        "status": batch_status,
        "batch_id": str(batch_id),
        "file_name": file_name,
        "sheet": sheet_name,
        "mapping": mapping,
        "rows_read": read,
        "rows_applied": applied,
        "rows_skipped": skipped,
        "rows_error": errors,
    }


def latest_yandex_plan_batch() -> dict[str, Any] | None:
    with engine.begin() as connection:
        row = connection.execute(
            text(
                """
                SELECT
                    id,
                    source_file_name,
                    source_file_modified,
                    source_file_size,
                    source_sheet,
                    status,
                    rows_read,
                    rows_applied,
                    rows_skipped,
                    rows_error,
                    message,
                    mapping,
                    created_at,
                    completed_at
                FROM erp_plan_import_batches
                ORDER BY created_at DESC
                LIMIT 1
                """
            )
        ).mappings().first()

    return dict(row) if row else None
